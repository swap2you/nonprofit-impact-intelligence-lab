"""Build the executive deck and a matching PDF from the final screenshots."""
import struct
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt
from reportlab.lib.pagesizes import landscape
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas

ROOT = Path(__file__).resolve().parents[1]
SHOTS = ROOT / "assets" / "screenshots"
OUT = ROOT / "assets" / "presentation"
PPTX = OUT / "Nonprofit_Impact_Intelligence_Lab_Executive_Demo.pptx"
PDF = OUT / "Nonprofit_Impact_Intelligence_Lab_Executive_Demo.pdf"

NAVY = RGBColor(0x10, 0x18, 0x20)
CARD = RGBColor(0x1B, 0x26, 0x36)
VIOLET = RGBColor(0x8B, 0x7C, 0xF7)
TEXT = RGBColor(0xE7, 0xEE, 0xF8)
MUTED = RGBColor(0x9A, 0xAB, 0xC0)
W, H = Inches(13.333), Inches(7.5)


def _set(run, text, size, color, bold=False):
    run.text = text
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.name = "Calibri"


def add_text(slide, text, x, y, w, h, size, color, bold=False):
    box = slide.shapes.add_textbox(x, y, w, h)
    frame = box.text_frame
    frame.word_wrap = True
    _set(frame.paragraphs[0], text, size, color, bold)
    return frame


def footer(slide, number, total=12):
    add_text(slide, "Synthetic demonstration  ·  not affiliated with buildOn", Inches(0.5), Inches(7.05), Inches(9), Inches(0.3), 12, MUTED)
    add_text(slide, f"{number}  /  {total}", Inches(11.4), Inches(7.05), Inches(1.4), Inches(0.3), 12, MUTED)


def blank(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = NAVY
    return slide


def card(slide, x, y, w, h):
    shape = slide.shapes.add_shape(1, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = CARD
    shape.line.fill.background()
    return shape


def build_pptx():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    prs.core_properties.title = "Nonprofit Impact Intelligence Lab"
    prs.core_properties.subject = "Synthetic executive demonstration"

    slide = blank(prs)
    add_text(slide, "SYNTHETIC DEMO", Inches(0.7), Inches(1.6), Inches(6), Inches(0.4), 16, VIOLET, True)
    add_text(slide, "Nonprofit Impact\nIntelligence Lab", Inches(0.7), Inches(2.1), Inches(8), Inches(2.2), 48, TEXT, True)
    add_text(slide, "Quality, funding lookup, migration readiness,\nand bounded analytics on one local dataset.", Inches(0.7), Inches(4.6), Inches(8), Inches(1.2), 22, MUTED)
    footer(slide, 1)

    slide = blank(prs)
    add_text(slide, "Why impact data gets hard to trust", Inches(0.6), Inches(0.4), Inches(12), Inches(0.7), 32, TEXT, True)
    points = [
        ("Late files", "Submissions arrive after the cutoff or sit stale."),
        ("Broken links", "A funding row can point at a site that does not exist."),
        ("Unclear movement", "A spike can be operations or a data defect."),
        ("Migration risk", "Source and target keys disagree more often than a slide admits."),
    ]
    for i, (title, body) in enumerate(points):
        x = Inches(0.6 + (i % 2) * 6.3)
        y = Inches(1.6 + (i // 2) * 2.4)
        card(slide, x, y, Inches(5.9), Inches(2.1))
        add_text(slide, title, x + Inches(0.3), y + Inches(0.35), Inches(5.2), Inches(0.5), 22, VIOLET, True)
        add_text(slide, body, x + Inches(0.3), y + Inches(1.0), Inches(5.2), Inches(0.7), 18, TEXT)
    footer(slide, 2)

    slide = blank(prs)
    add_text(slide, "What this lab demonstrates", Inches(0.6), Inches(0.4), Inches(12), Inches(0.7), 32, TEXT, True)
    items = [
        "Deterministic synthetic data, seed 20260923",
        "PostgreSQL 16 and a SQLite fallback",
        "Seven quality dimensions with a visible issue ledger",
        "FastAPI plus seven Streamlit pages",
        "CSV, XLSX, and PDF preset exports",
        "One-command Windows start and GitHub Actions",
    ]
    for i, item in enumerate(items):
        y = Inches(1.4 + i * 0.85)
        card(slide, Inches(0.6), y, Inches(12), Inches(0.7))
        add_text(slide, item, Inches(0.9), y + Inches(0.12), Inches(11), Inches(0.45), 20, TEXT)
    footer(slide, 3)

    slide = blank(prs)
    add_text(slide, "Architecture", Inches(0.6), Inches(0.4), Inches(12), Inches(0.7), 32, TEXT, True)
    steps = ["Generator", "PostgreSQL\nor SQLite", "FastAPI", "Streamlit\n7 pages", "Exports"]
    for i, label in enumerate(steps):
        x = Inches(0.55 + i * 2.55)
        card(slide, x, Inches(2.6), Inches(2.3), Inches(1.8))
        add_text(slide, label, x + Inches(0.1), Inches(3.05), Inches(2.1), Inches(1.0), 18, TEXT, True)
    add_text(slide, "Both the API and the dashboard read DATABASE_URL. There is no arbitrary SQL endpoint.", Inches(0.6), Inches(5.2), Inches(12), Inches(0.8), 18, MUTED)
    footer(slide, 4)

    slide = blank(prs)
    add_text(slide, "Synthetic scale", Inches(0.6), Inches(0.4), Inches(12), Inches(0.7), 32, TEXT, True)
    stats = [("8", "countries"), ("6", "programs"), ("240", "sites"), ("24", "months"), ("5,760", "enrollment rows"), ("23,745", "fact rows")]
    for i, (value, label) in enumerate(stats):
        x = Inches(0.5 + (i % 3) * 4.2)
        y = Inches(1.7 + (i // 3) * 2.4)
        card(slide, x, y, Inches(3.9), Inches(2.1))
        add_text(slide, value, x + Inches(0.25), y + Inches(0.3), Inches(3.4), Inches(0.9), 36, VIOLET, True)
        add_text(slide, label, x + Inches(0.25), y + Inches(1.25), Inches(3.4), Inches(0.5), 18, TEXT)
    footer(slide, 5)

    shots = [
        (6, "Executive Impact Overview", "01-executive-impact-overview.png", "Scale, freshness, quality score, and a rolling enrollment view."),
        (7, "Data Quality Command Center", "02-data-quality-command-center.png", "High, medium, and low are labeled in words. Record 97 is an orphan site reference."),
        (8, "Donor / Funding Lookup", "03-donor-funding-lookup.png", "A filtered briefing. Diagnostic demo, not a donor-ready certification."),
        (9, "Analytics and early warning", "06-analytics-early-warning.png", "Trend, IQR markers, and a one-period band. Not a causal claim."),
        (10, "Migration and reconciliation", "05-migration-reconciliation.png", "Readiness is about 10.8% because unresolved keys are planted."),
    ]
    for number, title, filename, caption in shots:
        slide = blank(prs)
        add_text(slide, title, Inches(0.45), Inches(0.25), Inches(8), Inches(0.5), 26, TEXT, True)
        add_text(slide, caption, Inches(0.45), Inches(0.8), Inches(12), Inches(0.4), 14, MUTED)
        with open(SHOTS / filename, "rb") as handle:
            handle.read(16)
            px_w, px_h = struct.unpack(">II", handle.read(8))
        max_w, max_h = Inches(12.4), Inches(5.15)
        scale = min(max_w / Emu(px_w * 9525), max_h / Emu(px_h * 9525))
        slide.shapes.add_picture(str(SHOTS / filename), Inches(0.45), Inches(1.35), int(px_w * 9525 * scale), int(px_h * 9525 * scale))
        footer(slide, number)

    slide = blank(prs)
    add_text(slide, "Run it, then prove it", Inches(0.6), Inches(0.4), Inches(12), Inches(0.7), 32, TEXT, True)
    card(slide, Inches(0.6), Inches(1.5), Inches(12), Inches(1.6))
    add_text(slide, ".\\scripts\\windows\\Start-ImpactLab.ps1", Inches(0.9), Inches(1.9), Inches(11), Inches(0.7), 28, VIOLET, True)
    add_text(slide, "Dashboard 127.0.0.1:8501    API docs 127.0.0.1:8000/docs", Inches(0.9), Inches(4.0), Inches(11), Inches(0.5), 20, TEXT)
    add_text(slide, "CI runs the same pytest suite on SQLite and on PostgreSQL 16.", Inches(0.9), Inches(4.7), Inches(11), Inches(0.5), 20, TEXT)
    footer(slide, 11)

    slide = blank(prs)
    add_text(slide, "Limits and a next step", Inches(0.6), Inches(0.4), Inches(12), Inches(0.7), 32, TEXT, True)
    limits = [
        "Synthetic only. No real donors, sites, or participants.",
        "Forecasts are illustrative ranges, not commitments.",
        "Low migration readiness is an intentional scenario.",
        "No production authentication or notification channels.",
    ]
    for i, line in enumerate(limits):
        add_text(slide, line, Inches(0.8), Inches(1.5 + i * 0.8), Inches(11), Inches(0.6), 22, TEXT)
    add_text(slide, "Extension path: configurable rules, authenticated database access, optional alerts.", Inches(0.8), Inches(5.2), Inches(11), Inches(0.7), 18, MUTED)
    footer(slide, 12)
    OUT.mkdir(parents=True, exist_ok=True)
    prs.save(PPTX)
    return PPTX


def build_pdf():
    page = landscape((960, 540))
    c = canvas.Canvas(str(PDF), pagesize=page)
    titles = [
        "Nonprofit Impact Intelligence Lab",
        "Why impact data gets hard to trust",
        "What this lab demonstrates",
        "Architecture",
        "Synthetic scale: 8 countries, 6 programs, 240 sites, 24 months, 23,745 fact rows",
        "Executive Impact Overview",
        "Data Quality Command Center",
        "Donor / Funding Lookup",
        "Analytics and early warning",
        "Migration and reconciliation",
        "Daily command: .\\scripts\\windows\\Start-ImpactLab.ps1",
        "Limits: synthetic data, illustrative forecast, no production claim",
    ]
    images = {
        5: "01-executive-impact-overview.png",
        6: "02-data-quality-command-center.png",
        7: "03-donor-funding-lookup.png",
        8: "06-analytics-early-warning.png",
        9: "05-migration-reconciliation.png",
    }
    for i, title in enumerate(titles):
        c.setFillColorRGB(0.063, 0.094, 0.125)
        c.rect(0, 0, 960, 540, fill=1, stroke=0)
        c.setFillColorRGB(0.91, 0.93, 0.97)
        c.setFont("Helvetica-Bold", 22)
        c.drawString(36, 490, title)
        c.setFillColorRGB(0.55, 0.49, 0.97)
        c.setFont("Helvetica", 11)
        c.drawString(36, 24, f"Synthetic demonstration   {i + 1} / 12")
        if i in images:
            c.drawImage(ImageReader(str(SHOTS / images[i])), 36, 50, width=888, height=420, preserveAspectRatio=True, anchor="c", mask="auto")
        c.showPage()
    c.save()


if __name__ == "__main__":
    print(build_pptx())
    build_pdf()
    print(PDF)
